"""Bulk import of ALREADY-TAKEN leave (Leave -> Bulk import taken leave,
Ganesh, 2026-09-15 — Norine had a spreadsheet of casual/sick leave dates
already taken this year and needed Planned Time balances to reflect them).

This is a different job from app/leave_bulk_upload.py's "Bulk assign
leaves" sheet, which only ever patches three annual entitlement NUMBERS on
an employee (Casual/Sick/Vacation Leaves/Year, display-only leftovers from
the pre-V2 leave system) — that sheet never creates a leave record and has
no connection to Leave Management V2's balances at all. This module does
the opposite: each row becomes a real, already-approved LeaveRecord, so
engine.leave_balance_v2()'s Used/Remaining for whichever type was entered
(Planned Time, Unplanned Time, ...) come out correct going forward.

Every row here is applied exactly as if an admin had used the "Record
leave" form on /admin/leave once per row (see leave_add() in
app/routes/admin.py) — same status=LEAVE_APPROVED,
approved_minutes_per_day=what-was-entered, requires_lead_review=False
reasoning (nobody requested it, so there's nothing for a Team Lead to
review). Every row must match an existing employee by Employee ID (the
same LOMK001-style employee_code every other bulk-upload sheet in this app
matches on) — this sheet never creates new employees.

Hours per day is optional: leave it blank for a full day off and the app
automatically uses that employee's own daily target for each day in
range (same "None means full day" convention as the single-entry form and
engine.leave_minutes_on) — more accurate than pre-computing hours by hand,
and correct for anyone on a non-standard schedule.

Re-uploading the same sheet (e.g. after fixing one row, or by accident) is
safe: process_upload() skips a row that exactly matches an
employee+date-range+type already on file, rather than creating a
duplicate leave record and double-counting it against that employee's
balance. There's no "update" mode for this sheet the way the entitlement
and holiday sheets have one, since a real historical event has no other
field to patch once it exists — deleting a wrong row is done from the
Leave Management page itself.

Row-parsing rules are plain functions, testable without a database, same
convention as every other *_bulk_upload.py module in this app.
"""
import datetime as dt
from typing import Dict, List, Optional, Set, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select

from app import models as m
from app.bulk_upload import parse_cell_date

MAX_ROWS = 1000

TEMPLATE_HEADERS = [
    "Employee ID", "Employee Name", "Start Date", "End Date", "Type",
    "Hours/day (blank = full day)", "Note",
]
COL_WIDTHS = [14, 24, 14, 14, 16, 26, 30]
COL_LETTERS = "ABCDEFG"


def parse_hours(raw) -> Optional[int]:
    """Blank -> None ("full day" — the app fills in that employee's own
    daily target for each date in range, same convention as the
    single-entry "Record leave" form's own Hours/day field). Otherwise a
    non-negative number of hours, converted to whole minutes (rounded)."""
    if raw is None or str(raw).strip() == "":
        return None
    try:
        hours = float(raw)
    except (TypeError, ValueError):
        raise ValueError("Hours/day must be a number, e.g. 4 or 2.5 (blank = full day)")
    if hours < 0:
        raise ValueError("Hours/day can't be negative")
    return int(round(hours * 60))


def parse_row(raw: dict, code_to_id: Dict[str, int], valid_types: Tuple[str, ...]) -> dict:
    """Returns one of:
      {"mode": "ok", "employee_id": int, "start": date, "end": date,
       "type": str, "minutes": Optional[int], "note": str, "error": None}
      {"mode": "error", "error": "..."}
    "type" in the "ok" result is always one of valid_types (matched
    case-insensitively against what was typed in the sheet, then
    normalized to the canonical spelling)."""
    emp_id_raw = str(raw.get("Employee ID") or "").strip()
    if not emp_id_raw:
        return {"mode": "error", "error": "Employee ID is required — this sheet never creates new employees"}
    employee_id = code_to_id.get(emp_id_raw.upper())
    if employee_id is None:
        return {"mode": "error", "error": f"Employee ID '{emp_id_raw}' not found"}

    try:
        start = parse_cell_date(raw.get("Start Date"))
    except ValueError as e:
        return {"mode": "error", "error": str(e)}
    if start is None:
        return {"mode": "error", "error": "Start Date is required"}

    end_raw = raw.get("End Date")
    if end_raw is None or str(end_raw).strip() == "":
        end = start
    else:
        try:
            end = parse_cell_date(end_raw)
        except ValueError as e:
            return {"mode": "error", "error": str(e)}
        if end is None:
            end = start
    if end < start:
        return {"mode": "error", "error": "End Date is before Start Date"}

    leave_type_raw = str(raw.get("Type") or "").strip()
    if not leave_type_raw:
        return {"mode": "error", "error": "Type is required"}
    matched_type = next((t for t in valid_types if t.lower() == leave_type_raw.lower()), None)
    if matched_type is None:
        return {"mode": "error", "error": f"'{leave_type_raw}' isn't a valid leave type — use one of: {', '.join(valid_types)}"}

    try:
        minutes = parse_hours(raw.get("Hours/day (blank = full day)"))
    except ValueError as e:
        return {"mode": "error", "error": str(e)}

    note = str(raw.get("Note") or "").strip()

    return {
        "mode": "ok", "employee_id": employee_id, "start": start, "end": end,
        "type": matched_type, "minutes": minutes, "note": note, "error": None,
    }


def read_upload_rows(wb: Workbook) -> Tuple[List[dict], Optional[str]]:
    """Same header-mapping shape as every other *_bulk_upload.py module in
    this app — see e.g. app/holiday_bulk_upload.py's own copy of this."""
    ws = wb.active
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None or all(c is None or str(c).strip() == "" for c in header_row):
        return [], "The sheet is empty — no header row found."
    header_map = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        header_map[str(cell).strip().lower()] = idx
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue  # blank row — skip silently, not an error
        rows.append({
            field: (row[header_map[field.lower()]]
                     if field.lower() in header_map and header_map[field.lower()] < len(row)
                     else None)
            for field in TEMPLATE_HEADERS
        })
    return rows, None


def process_upload(db, wb: Workbook, valid_types: Tuple[str, ...], entered_by: str) -> dict:
    """Parses + applies an uploaded workbook. Every valid, non-duplicate row
    becomes one already-approved LeaveRecord, committed in one transaction.
    A row whose employee+start+end+type exactly matches a LeaveRecord
    already on file is skipped as a duplicate rather than double-counted —
    see module docstring for why re-uploading is safe. Invalid rows are
    skipped and listed with a reason, never silently dropped. Returns
    {"added": int,
     "recompute_ranges": {employee_id: (min_start, max_end)},
     "skipped": [{"row": int, "name": str, "reason": str}],
     "header_error": str | None}."""
    rows, header_error = read_upload_rows(wb)
    if header_error:
        return {"added": 0, "recompute_ranges": {}, "skipped": [], "header_error": header_error}
    if len(rows) > MAX_ROWS:
        return {
            "added": 0, "recompute_ranges": {}, "skipped": [],
            "header_error": f"Sheet has {len(rows)} data rows — max is {MAX_ROWS} per upload. Split it into batches.",
        }

    existing_emps = list(db.execute(select(m.Employee.id, m.Employee.employee_code)).all())
    code_to_id = {code.upper(): eid for (eid, code) in existing_emps if code}

    existing_keys: Set[Tuple[int, dt.date, dt.date, str]] = set(
        db.execute(
            select(m.LeaveRecord.employee_id, m.LeaveRecord.start_date, m.LeaveRecord.end_date, m.LeaveRecord.type)
        ).all()
    )

    added = 0
    recompute_ranges: Dict[int, Tuple[dt.date, dt.date]] = {}
    skipped = []
    for i, raw in enumerate(rows, start=2):  # row 1 is the header
        display = (str(raw.get("Employee Name") or raw.get("Employee ID") or "").strip()) or "(blank)"
        result = parse_row(raw, code_to_id, valid_types)
        if result["error"]:
            skipped.append({"row": i, "name": display, "reason": result["error"]})
            continue
        key = (result["employee_id"], result["start"], result["end"], result["type"])
        if key in existing_keys:
            skipped.append({
                "row": i, "name": display,
                "reason": f"Already recorded ({result['type']}, {result['start']}–{result['end']}) — skipped as a duplicate",
            })
            continue
        existing_keys.add(key)  # a duplicate row within this same file is also caught
        lv = m.LeaveRecord(
            employee_id=result["employee_id"], start_date=result["start"], end_date=result["end"],
            type=result["type"], minutes_per_day=result["minutes"], note=result["note"],
            entered_by=entered_by, status=m.LEAVE_APPROVED,
            approved_minutes_per_day=result["minutes"], requires_lead_review=False,
        )
        db.add(lv)
        added += 1
        prev = recompute_ranges.get(result["employee_id"])
        if prev is None:
            recompute_ranges[result["employee_id"]] = (result["start"], result["end"])
        else:
            recompute_ranges[result["employee_id"]] = (min(prev[0], result["start"]), max(prev[1], result["end"]))

    if added:
        db.commit()
    return {"added": added, "recompute_ranges": recompute_ranges, "skipped": skipped, "header_error": None}


def build_sample_workbook(valid_types: Tuple[str, ...]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leave taken"
    ws.append(TEMPLATE_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    sample_type = valid_types[0] if valid_types else "Planned Time"
    today = dt.date.today()
    ws.append(["LOMK001", "Jane Doe", today, today, sample_type, "", "Imported from HR spreadsheet"])
    for col, width in zip(COL_LETTERS, COL_WIDTHS):
        ws.column_dimensions[col].width = width

    info = wb.create_sheet("Instructions")
    info.append(["Column", "Required?", "Format / allowed values"])
    for c in info[1]:
        c.font = Font(bold=True)
    for row in [
        ("Employee ID", "Yes", "Must match an existing employee's ID (LOMK001, ...) — this sheet never creates new employees"),
        ("Employee Name", "No", "For your own reference only — not used to match the row, Employee ID is"),
        ("Start Date", "Yes", "YYYY-MM-DD, or an Excel date cell"),
        ("End Date", "No", "Blank = same as Start Date (one day). Must not be before Start Date"),
        ("Type", "Yes", "One of: " + ", ".join(valid_types)),
        ("Hours/day (blank = full day)", "No", "A number of hours, e.g. 4 or 2.5. Blank = full day (uses that employee's own daily target)"),
        ("Note", "No", "Free text, e.g. where this record came from"),
        ("", "", ""),
        ("Each row creates ONE already-approved leave record, exactly like", "", ""),
        ("using \"Record leave\" on the Leave page once per row.", "", ""),
        ("Re-uploading the same sheet is safe — a row that exactly matches", "", ""),
        ("an employee+date range+type already on file is skipped instead", "", ""),
        ("of creating a duplicate, so fixing one row and re-uploading is fine.", "", ""),
    ]:
        info.append(row)
    info.column_dimensions["A"].width = 55
    info.column_dimensions["B"].width = 12
    info.column_dimensions["C"].width = 75
    return wb
